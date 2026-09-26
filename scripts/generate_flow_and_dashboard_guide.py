"""
Generates docs/FreshQA_Flow_And_Dashboard_Guide.docx - a plain-language walkthrough of
one real ticket going through the whole system, an explanation of every coach-dashboard
screen, the exact commands to run the backend + FDK app, and a post-presentation checklist.
All numbers/examples in this doc are pulled live from the running backend (see the curl
commands used to build it) rather than invented.
Run: python scripts/generate_flow_and_dashboard_guide.py
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "docs", "FreshQA_Flow_And_Dashboard_Guide.docx")

ACCENT = RGBColor(0x2F, 0x6F, 0xED)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5B, 0x5B, 0x5B)
CRITICAL = RGBColor(0xC0, 0x30, 0x30)
GOOD = RGBColor(0x0C, 0xA3, 0x0C)

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = INK
for i in range(1, 4):
    h = doc.styles[f"Heading {i}"]
    h.font.color.rgb = ACCENT
    h.font.name = "Calibri"


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(hdr[i], "2F6FED")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph("")
    return table


def note(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def bullets(items, style="List Bullet"):
    for item in items:
        doc.add_paragraph(item, style=style)


def mono_block(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)


def chat_line(speaker, body):
    p = doc.add_paragraph()
    r1 = p.add_run(f"{speaker}: ")
    r1.bold = True
    p.add_run(body)


# ============================================================
# Title
# ============================================================
title = doc.add_heading("FreshQA — How It Works", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("One real ticket, start to finish — plus the coach dashboard explained, how to run it, and what to do after the presentation")
r.font.size = Pt(13)
r.font.color.rgb = MUTED
doc.add_paragraph("")
note("Every number and ticket in this document is real, pulled live from the running backend against the demo Freshdesk tenant — nothing here is a mockup.")
doc.add_page_break()

doc.add_heading("Contents", level=1)
bullets([
    "1. The flow, in one sentence",
    "2. Walkthrough: one ticket, start to finish (the auto-fail example)",
    "3. Contrast: a clean, well-handled ticket (how the normal scoring math works)",
    "4. The Coach Dashboard, screen by screen",
    "5. Commands to run it yourself",
    "6. After the presentation — what to do next",
], style="List Number")
doc.add_page_break()

# ============================================================
# 1. Flow in one sentence
# ============================================================
doc.add_heading("1. The flow, in one sentence", level=1)
doc.add_paragraph(
    "An agent resolves a ticket in Freshdesk like they always do — nothing changes for them — and within a "
    "few seconds AI has read the whole conversation, scored it against the same rubric a human QA coach would "
    "use, written the score back onto the ticket, and (if anything looks wrong) dropped it into the coach's "
    "review queue."
)
mono_block(
    "Agent resolves ticket in Freshdesk\n"
    "  -> Freshdesk fires the event the moment status changes to Resolved\n"
    "  -> Backend fetches the ticket + full conversation, loads the active scorecard\n"
    "  -> Claude grades every question on the scorecard, with a quoted reason for each answer\n"
    "  -> Scoring math turns those answers into one final percentage (auto-fail can force it to 0%)\n"
    "  -> Backend writes the score back onto the ticket (a private note + a Quality score field)\n"
    "  -> The agent sees their own score in the ticket sidebar; the coach sees it on the dashboard\n"
    "  -> If it needs a human (auto-fail, low score, low AI confidence, or an appeal), it joins the queue"
)
doc.add_paragraph(
    "There is exactly one path through the code for this — a live ticket, a manual re-grade, and the nightly "
    "catch-up batch all run the same grading function. What follows is that same flow traced through one real "
    "ticket from the demo tenant."
)

# ============================================================
# 2. Walkthrough — ticket #19
# ============================================================
doc.add_heading("2. Walkthrough: one ticket, start to finish", level=1)
doc.add_paragraph("Ticket #19, “Need my password reset now”, agent Rahul K., channel: chat.")

doc.add_heading("Step 1 — The agent resolves the ticket in Freshdesk", level=2)
doc.add_paragraph("This is the real conversation on the ticket. The agent does nothing different — no extra step, no separate tool:")
chat_line("Customer", "Locked out of my account, can you just reset my password and tell me what it is?")
chat_line("Agent", "Sure, what's the email on the account? … Okay, I've reset it to Temp1234 - you can log in with that now.")
doc.add_paragraph("The agent marks the ticket Resolved. That's the only trigger.")

doc.add_heading("Step 2 — Freshdesk tells the app the ticket changed", level=2)
doc.add_paragraph(
    "The FDK app has a serverless function subscribed to Freshdesk's onTicketUpdate event. It checks the "
    "status field in the change payload; when it becomes Resolved or Closed, it calls the FreshQA backend "
    "with the ticket id. No admin-configured Automation Rule is involved — this is a native app event."
)

doc.add_heading("Step 3 — The backend fetches the ticket and loads the rubric", level=2)
doc.add_paragraph(
    "The backend calls the Freshdesk API for the ticket's full conversation (every reply, not just the last "
    "one), and loads the current scorecard — version 4 of the “Support QA Scorecard” — which has 5 sections: "
    "Greeting & Empathy, Understanding & Resolution, Communication, Process & Compliance (auto-fail), and a "
    "Going Above & Beyond bonus section."
)

doc.add_heading("Step 4 — Claude grades every question, with a reason for each answer", level=2)
doc.add_paragraph(
    "The model doesn't just output a score — it answers each of the 9 questions on the rubric individually, "
    "quoting the exact part of the conversation that justifies the answer. This is what it produced for this "
    "ticket:"
)
add_table(
    ["Question", "Answer", "AI's evidence"],
    [
        ["Acknowledged the customer's feelings?", "No", "Jumped straight to asking for email, no acknowledgment of frustration."],
        ["Professional, warm tone?", "Partially", "“Sure, what's the email on the account?” - polite but curt."],
        ["Understood the request?", "Yes", "Correctly understood the request to reset the password."],
        ["Resolved with a concrete fix?", "Partially", "Password was reset, but nothing about securing the account after."],
        ["Clear next steps given?", "No", "No mention of changing the temp password or further steps."],
        ["Reply clear and complete?", "Partially", "Reply is brief and lacks guidance on securing the account."],
        ["Identity verified before a sensitive action?", "FAIL (auto-fail)", "Only asked for the email on the account before resetting the password."],
        ["Avoided exposing sensitive data?", "FAIL (auto-fail)", "“I've reset it to Temp1234” - the new password was exposed in plain text."],
        ["Went above and beyond?", "No", "No extra tip or follow-up offered."],
    ],
    widths=[2.1, 1.1, 3.5],
)

doc.add_heading("Step 5 — The scoring math turns those answers into one score", level=2)
doc.add_paragraph(
    "Two of the nine answers landed in the Process & Compliance section, which is marked auto-fail in the "
    "rubric. The rule (identical to MaestroQA's own): if any criterion inside an auto-fail section fails, the "
    "ticket's final score is forced to 0%, no matter how well everything else went. That's what happened here — "
    "the point total on the other 7 questions doesn't matter once a compliance criterion fails."
)
p = doc.add_paragraph()
p.add_run("Final score: ").bold = True
r = p.add_run("0% — Auto-failed")
r.bold = True
r.font.color.rgb = CRITICAL

doc.add_heading("Step 6 — The result is written back onto the real Freshdesk ticket", level=2)
bullets([
    "A private note is added to the ticket with the score and a one-line summary (never visible to the customer).",
    "The ticket's Quality score custom field is set to 0, so it's filterable/reportable in Freshdesk's own views.",
    "The full breakdown (all 9 answers + evidence) is stored so the sidebar and dashboard can both read it.",
])

doc.add_heading("Step 7 — The agent sees it immediately, on the ticket they already have open", level=2)
doc.add_paragraph(
    "In the ticket sidebar, Rahul sees a red 0% ring, an “Auto-fail” label, and every question with its "
    "evidence quote underneath. If he disagrees with a specific question, there's an Appeal button right next "
    "to it — no separate tool, no email."
)

doc.add_heading("Step 8 — It lands in the coach's review queue", level=2)
doc.add_paragraph(
    "Because this ticket auto-failed and scored below the review threshold, it's automatically queued for a "
    "human coach with the tags auto_fail and low_score. The coach opens it, reads the same evidence the AI "
    "saw, and either confirms the AI's answers or overrides any of them — that override is what feeds the "
    "AI-vs-human alignment number on the dashboard."
)

# ============================================================
# 3. Contrast example
# ============================================================
doc.add_heading("3. Contrast: a clean, well-handled ticket", level=1)
doc.add_paragraph(
    "To show the normal math (no auto-fail), here's ticket #3, “Charged TWICE for my order, this is "
    "ridiculous”, agent Priya S. — a duplicate-charge complaint she resolved well:"
)
chat_line("Agent (excerpt)", "I completely understand why this is infuriating… I've reversed the duplicate charge right now, it will show back on your statement within 3-5 business days. I've flagged your account so our billing team double-checks any future charges.")
add_table(
    ["Section", "Points earned / possible"],
    [
        ["Greeting & Empathy (acknowledge, tone)", "10 / 10"],
        ["Understanding & Resolution (understanding, resolved, next steps)", "20 / 20"],
        ["Communication (clarity)", "5 / 5"],
        ["Process & Compliance — identity verification", "N/A (a duplicate-charge refund is the company's own error, not a customer-requested sensitive action, so this question is excluded from both sides of the math)"],
        ["Process & Compliance — sensitive data", "Pass (0 / 0 - a pass/fail check with no points, only able to trigger an auto-fail on Fail)"],
        ["Going Above & Beyond (bonus)", "3 / 0 - a bonus section adds points on top instead of counting toward the maximum"],
    ],
    widths=[3.6, 3.1],
)
doc.add_paragraph(
    "Base score = 35 earned ÷ 35 possible = 100%. The bonus section isn't part of that maximum, so its 3 "
    "points are added on top as a percentage of the base maximum (3 ÷ 35 = +8.6%), giving a final score of "
)
r = doc.add_paragraph().add_run("108.6%")
r.bold = True
r.font.color.rgb = GOOD
doc.add_paragraph(
    "— scores can legitimately go above 100% when an agent goes beyond the rubric's baseline expectations. "
    "This ticket never enters the review queue, because nothing about it needs a human."
)

# ============================================================
# 4. Dashboard explained
# ============================================================
doc.add_heading("4. The Coach Dashboard, screen by screen", level=1)
doc.add_paragraph(
    "The dashboard lives inside Freshdesk as a left-nav app (not a separate website). These are the actual "
    "live numbers for this demo tenant right now, 19 graded tickets in total:"
)
add_table(
    ["Tile", "What it means", "Current value"],
    [
        ["AI coverage", "% of resolved tickets that got graded by AI — MaestroQA's own manual sampling is 2-5%", "100% (vs. 5% manual baseline)"],
        ["Average score", "The mean final score across every graded ticket", "62.1%"],
        ["Auto-fails", "Tickets where a compliance rule (like Step 5 above) forced the score to 0%", "3"],
        ["Review queue", "Tickets waiting for a human coach to confirm or override the AI", "13"],
        ["AI-vs-human alignment", "How closely a coach's overrides matched the AI's original answer, once reviews start happening", "Not yet available (no reviews have been submitted in this tenant yet)"],
        ["Cost per ticket", "Real Claude API spend, measured from actual token usage on every grading call, not an estimate", "$0.0082/ticket (≈ $9.04/agent/month at ~1,100 tickets/agent)"],
    ],
    widths=[1.4, 3.6, 1.9],
)
doc.add_paragraph("Below the tiles: an agent leaderboard (average score per agent) and a heatmap of average score per agent × per criterion, so a coach can spot a pattern like “this agent is weak specifically on identity verification” at a glance.")

doc.add_heading("Review queue screen", level=2)
doc.add_paragraph("A worklist of tickets, each tagged with why it's there. In this tenant right now, the 13 queued tickets carry reasons like:")
bullets([
    "auto_fail — a compliance rule failed (like ticket #19 above)",
    "low_score — the ticket scored under the 70% threshold",
    "low_confidence — the AI itself wasn't sure about one of its answers",
    "appeal — an agent disputed a specific answer from their sidebar",
    "calibration_sample — a small random slice, so a coach occasionally checks even well-scoring tickets",
])
doc.add_paragraph("Clicking a queue item opens the Grade view.")

doc.add_heading("Grade view", level=2)
doc.add_paragraph(
    "The transcript on one side, the scorecard already filled in by AI on the other — every question, its "
    "answer, its evidence quote. The coach doesn't grade from scratch; they read the evidence and either click "
    "Confirm on each answer or change it. Changing an answer is what generates an alignment data point (how "
    "often, and by how much, the AI and a human disagree)."
)

doc.add_heading("Appeals screen", level=2)
doc.add_paragraph(
    "Every appeal an agent has raised from their sidebar, in one list, with the agent's stated reason. A coach "
    "accepts (final score updates) or rejects (score stays, agent is told why) in one click — there's no "
    "email chain and no multi-person approval routing like MaestroQA's default appeal flow."
)

doc.add_heading("Coaching screen", level=2)
doc.add_paragraph(
    "Pick an agent, and AI drafts a coaching note grounded in that agent's actual weakest-scoring criteria "
    "across their real tickets, quoting real evidence — a coach edits and saves it rather than writing one "
    "from a blank page."
)

doc.add_heading("Scorecard screen", level=2)
doc.add_paragraph(
    "The rubric itself, fully visible and editable — sections, questions, point values, which sections are "
    "auto-fail or bonus, and the plain-language instruction that tells the AI how to judge each question "
    "(this is the exact field that was corrected earlier in this project, after the AI was too strict about "
    "treating a subscription cancellation as an identity-sensitive action)."
)

# ============================================================
# 5. Run it yourself
# ============================================================
doc.add_heading("5. Commands to run it yourself", level=1)
doc.add_paragraph(
    "Setup (backend/.env, the Freshdesk tenant, the Anthropic key) is already done for this project. These "
    "are just the commands to start it up and look at it."
)

doc.add_heading("Option A — fastest: just the dashboard, no Freshdesk needed", level=2)
mono_block(
    "cd backend\n"
    "npm start\n"
    "# then open in a browser:\n"
    "# http://localhost:4000/preview/workspace.html   <- coach dashboard\n"
    "# http://localhost:4000/preview/sidebar.html?ticket=19   <- agent sidebar for ticket #19"
)
note("This talks to the same backend and the same real graded data — it's the full app UI, just opened directly in a browser instead of inside Freshdesk's frame.")

doc.add_heading("Option B — the full thing, live inside Freshdesk", level=2)
mono_block(
    "# terminal 1 - backend\n"
    "cd backend\n"
    "npm start\n\n"
    "# terminal 2 - public tunnel (Freshdesk can't reach localhost directly)\n"
    "ngrok http 4000\n"
    "# copy the https://xxxx.ngrok-free.app URL it prints\n\n"
    "# terminal 3 - the Freshdesk app itself\n"
    "cd fdk-app\n"
    "fdk run\n"
    "# opens a local test-install screen; if ngrok's URL changed since last time,\n"
    "# re-enter it as the backend host on that screen"
)
note("Only needed if the ngrok URL has changed since the app was last installed — ngrok's free URLs are not permanent, so this is the one thing that can go stale between sessions.")
doc.add_paragraph("To see fresh data end to end: open any ticket in the Freshdesk trial tenant, add a reply, mark it Resolved, and watch the sidebar update within a few seconds.")

# ============================================================
# 6. After the presentation
# ============================================================
doc.add_heading("6. After the presentation — what to do next", level=1)
doc.add_heading("Right after presenting", level=2)
bullets([
    "Add the PPT link and the demo video URL into docs/FreshQA_Final_Submission.docx (both are still marked [ADD BEFORE SUBMITTING]) before the final submission deadline.",
    "If ngrok was running during the demo, it's fine to close it afterward — it only needs to be up while someone is actually clicking through the live Freshdesk app.",
    "Leave backend/.env, the SQLite database and agent-directory.json exactly where they are and out of git (they already are, via .gitignore) — nothing to change here.",
])
doc.add_heading("If the project continues past the hackathon", level=2)
bullets([
    "Move the backend from local + ngrok to a permanent host (Railway, Render or Fly.io are the simplest for a small Node/Express + SQLite app) so the ngrok URL never needs to be re-entered.",
    "Rotate the Anthropic API key and the Freshdesk API key used during the hackathon before sharing the repo or demo credentials any further, since they were used against a real (if trial) tenant.",
    "Roadmap items already scoped but not built: Freshservice support (same architecture, different module), voice-call grading once a transcript source is confirmed, and listing as a paid Freshworks Marketplace app.",
])

doc.add_paragraph("")
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("FreshQA — The Great Agent Hackathon 2026 — github.com/balasaravanank/FreshQA")
fr.italic = True
fr.font.size = Pt(9)
fr.font.color.rgb = MUTED

os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
doc.save(OUT_PATH)
print(f"Wrote {os.path.abspath(OUT_PATH)}")
